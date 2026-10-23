
import os

os.environ["PADDLE_PDX_ENABLE_MKLDNN_BYDEFAULT"] = "False"

import re
import shutil
import tempfile
import traceback
from pathlib import Path

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from starlette.background import BackgroundTask
from paddleocr import PPStructureV3
from docx import Document
from docx.shared import Inches, Pt
import html
from bs4 import BeautifulSoup

app = FastAPI(
    title="Img2Word OCR API",
    description="Convert images into editable Word documents",
    version="2.0.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Initialize the document-aware OCR pipeline once at startup.
ocr = PPStructureV3(
    lang="en",
    use_doc_orientation_classify=True,
    use_doc_unwarping=True,
    use_textline_orientation=True,
)


def clean_inline_markdown(text):
    """Remove common inline Markdown syntax while keeping its text."""
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"__(.*?)__", r"\1", text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
    text = re.sub(r"(?<!_)_([^_]+)_(?!_)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def split_table_row(line):
    """Split a Markdown table row into its cell values."""
    line = line.strip()

    if line.startswith("|"):
        line = line[1:]

    if line.endswith("|"):
        line = line[:-1]

    return [cell.strip() for cell in line.split("|")]


def is_table_separator(line):
    """Identify Markdown table separator rows."""
    cells = split_table_row(line)

    return bool(cells) and all(
        re.fullmatch(r":?-{3,}:?", cell.replace(" ", ""))
        for cell in cells
    )


def add_markdown_table(document, table_lines):
    """Create a native, editable Word table from Markdown."""
    rows = []

    for line in table_lines:
        if is_table_separator(line):
            continue

        rows.append(split_table_row(line))

    if not rows:
        return

    column_count = max(len(row) for row in rows)

    if column_count == 0:
        return

    table = document.add_table(rows=1, cols=column_count)
    table.style = "Table Grid"
    table.autofit = True

    for row_index, values in enumerate(rows):
        if row_index == 0:
            cells = table.rows[0].cells
        else:
            cells = table.add_row().cells

        for column_index in range(column_count):
            value = values[column_index] if column_index < len(values) else ""
            cells[column_index].text = clean_inline_markdown(value)

            if row_index == 0:
                for run in cells[column_index].paragraphs[0].runs:
                    run.bold = True


def convert_html_tables_to_markdown(text):
    """Convert HTML tables to Markdown and remove unwanted HTML markup."""

    table_pattern = re.compile(
        r"<table\b[^>]*>.*?</table\s*>",
        flags=re.IGNORECASE | re.DOTALL,
    )

    def replace_table(match):
        soup = BeautifulSoup(match.group(0), "html.parser")
        table = soup.find("table")

        if table is None:
            return ""

        rows = []

        for row in table.find_all("tr"):
            cells = row.find_all(["th", "td"], recursive=False)

            if not cells:
                continue

            cell_values = []

            for cell in cells:
                value = cell.get_text(" ", strip=True)
                value = re.sub(r"\s+", " ", value)
                value = value.replace("|", r"\|")
                cell_values.append(value)

            rows.append(cell_values)

        if not rows:
            return ""

        column_count = max(len(row) for row in rows)
        normalized_rows = [
            row + [""] * (column_count - len(row))
            for row in rows
        ]

        header = normalized_rows[0]
        separator = ["---"] * column_count

        markdown_rows = [
            "| " + " | ".join(header) + " |",
            "| " + " | ".join(separator) + " |",
        ]

        for row in normalized_rows[1:]:
            markdown_rows.append(
                "| " + " | ".join(row) + " |"
            )

        return "\n\n" + "\n".join(markdown_rows) + "\n\n"

    # Convert tables before removing the remaining HTML tags.
    text = table_pattern.sub(replace_table, text)

    # Remove image markup so HTML image references do not appear as text.
    text = re.sub(
        r"<img\b[^>]*>",
        "",
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )

    # Remove remaining HTML tags but retain their textual content.
    text = re.sub(r"<[^>]+>", " ", text)
    text = html.unescape(text)

    # Clean excessive whitespace without merging separate lines.
    cleaned_lines = [
        re.sub(r"[ \t]+", " ", line).strip()
        for line in text.splitlines()
    ]

    return "\n".join(cleaned_lines)



def markdown_to_word(markdown_text, output_path):
    """Convert Markdown headings, paragraphs, lists, and tables to Word."""
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)

    lines = markdown_text.splitlines()
    index = 0

    while index < len(lines):
        line = lines[index].strip()

        if not line:
            index += 1
            continue

        # Markdown table
        if (
            "|" in line
            and index + 1 < len(lines)
            and is_table_separator(lines[index + 1])
        ):
            table_lines = [line]
            index += 1

            while index < len(lines) and "|" in lines[index]:
                table_lines.append(lines[index].strip())
                index += 1

            add_markdown_table(document, table_lines)
            document.add_paragraph()
            continue

        # Headings
        heading_match = re.match(r"^(#{1,6})\s+(.+)$", line)

        if heading_match:
            level = min(len(heading_match.group(1)), 3)
            heading_text = clean_inline_markdown(heading_match.group(2))
            document.add_heading(heading_text, level=level)
            index += 1
            continue

        # Bullet lists
        bullet_match = re.match(r"^[-*+]\s+(.+)$", line)

        if bullet_match:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(clean_inline_markdown(bullet_match.group(1)))
            index += 1
            continue

        # Numbered lists
        number_match = re.match(r"^\d+[.)]\s+(.+)$", line)

        if number_match:
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run(clean_inline_markdown(number_match.group(1)))
            index += 1
            continue

        # Horizontal rules and image-only Markdown lines
        if re.fullmatch(r"[-*_]{3,}", line):
            index += 1
            continue

        if re.fullmatch(r"!\[[^\]]*\]\([^)]+\)", line):
            index += 1
            continue

        # Ordinary text paragraph
        paragraph_lines = [line]
        index += 1

        while index < len(lines):
            next_line = lines[index].strip()

            if not next_line:
                break

            if re.match(r"^#{1,6}\s+", next_line):
                break

            if re.match(r"^[-*+]\s+", next_line):
                break

            if re.match(r"^\d+[.)]\s+", next_line):
                break

            if (
                "|" in next_line
                and index + 1 < len(lines)
                and is_table_separator(lines[index + 1])
            ):
                break

            paragraph_lines.append(next_line)
            index += 1

        paragraph_text = clean_inline_markdown(" ".join(paragraph_lines))

        if paragraph_text:
            document.add_paragraph(paragraph_text)

    if not document.paragraphs and not document.tables:
        document.add_paragraph("No readable text was detected in this image.")

    document.save(output_path)


@app.get("/")
def root():
    return {
        "message": "Img2Word OCR API is running",
        "status": "healthy",
    }


@app.get("/health")
def health_check():
    return {"status": "ok"}


@app.post("/convert")
async def convert_image(file: UploadFile = File(...)):
    allowed_extensions = {
        ".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"
    }

    extension = Path(file.filename or "").suffix.lower()

    if extension not in allowed_extensions:
        raise HTTPException(
            status_code=400,
            detail="Unsupported image format. Upload PNG, JPG, BMP, or TIFF.",
        )

    image_bytes = await file.read()

    if not image_bytes:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    if len(image_bytes) > 15 * 1024 * 1024:
        raise HTTPException(
            status_code=413,
            detail="Image size must be 15 MB or less.",
        )

    temp_dir = tempfile.mkdtemp()

    try:
        image_path = os.path.join(temp_dir, f"input{extension}")
        output_path = os.path.join(temp_dir, "converted.docx")
        markdown_dir = os.path.join(temp_dir, "markdown")
        os.makedirs(markdown_dir, exist_ok=True)

        with open(image_path, "wb") as image_file:
            image_file.write(image_bytes)

        # Detect document layout and reconstruct reading order.
        results = ocr.predict(image_path)

        markdown_pages = []

        for page_index, page_result in enumerate(results):
            page_dir = os.path.join(
                markdown_dir,
                f"page_{page_index + 1}",
            )
            os.makedirs(page_dir, exist_ok=True)

            # Use PaddleOCR's supported Markdown export.
            page_result.save_to_markdown(save_path=page_dir)

            markdown_files = sorted(Path(page_dir).rglob("*.md"))

            for markdown_file in markdown_files:
                markdown_text = markdown_file.read_text(encoding="utf-8").strip()

                if markdown_text:
                    markdown_pages.append(markdown_text)

        if not markdown_pages:
            raise RuntimeError(
                "The OCR pipeline did not produce Markdown output."
            )

        combined_markdown = convert_html_tables_to_markdown(
            "\n\n".join(markdown_pages)
        )

        markdown_to_word(combined_markdown, output_path)

        return FileResponse(
            output_path,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            filename="Img2Word_Converted.docx",
            background=BackgroundTask(
                shutil.rmtree,
                temp_dir,
                ignore_errors=True,
            ),
        )

    except HTTPException:
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise

    except Exception as error:
        shutil.rmtree(temp_dir, ignore_errors=True)
        traceback.print_exc()

        raise HTTPException(
            status_code=500,
            detail=f"Image conversion failed: {type(error).__name__}: {error}",
        ) from error
